"""THE ADTINI SLIDES, BUILT FROM THE QUOTE (2026-09-19, Kiri).

The proposal document is Brendan's letter; this is the deck the partner
presents. Same quote, five slides, and every one of them is conditional on
what the quote actually carries:

  1. the product and the SERP that was captured for it,
  2. the strategies that were sold -- and only those,
  3. the keyword list and its tier split, when SEO or AI Search is on it,
  4. the audit, when the audit is on it,
  5. the three tiers, with the AI Search line and the add-on market box
     printed only when the quote carries them.

A slide with nothing behind it is not printed at all, rather than printed
empty: a deck that shows a Website Audit section to a client who did not buy
one is a deck that has to be edited by hand before it can be sent, which is
the thing this is meant to stop.
"""
import io
import os

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

# ---------------------------------------------------------------- brand
NAVY = RGBColor(0x12, 0x3A, 0x63)
NAVY_D = RGBColor(0x0E, 0x2E, 0x4F)
BLUE = RGBColor(0x1C, 0x5B, 0xC4)
BLUE_MID = RGBColor(0x1F, 0x5F, 0xAD)
BLUE_LT = RGBColor(0x5E, 0x9B, 0xD6)
GREEN = RGBColor(0x2F, 0xA8, 0x4F)
LINE = RGBColor(0xDD, 0xE1, 0xE7)
MUTED = RGBColor(0x66, 0x70, 0x84)
INK = RGBColor(0x1A, 0x23, 0x30)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
# ON A BLUE CARD, GREY IS NEARLY GONE. The tiers' own screens print the
# carried-over lines a shade back from white, not in the muted grey the white
# card uses -- on the blue they came out unreadable.
TINT = RGBColor(0xDA, 0xE6, 0xF6)
FONT = "Poppins"

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

TITLE = "Search Engine Optimization+"

# THE DECK'S OWN WORDS, off the adtini product slides. Nothing here is
# generated: the client reads the same sentences whoever built the quote.
COPY = {
    "tagline": "The best long-term strategy for driving online leads, "
               "improving a website’s search engine visibility and "
               "keyword rankings.",
    "include_heading": "Your SEO Campaign Will Include:",
    "include": [
        ("Keyword Targeting:",
         " We analyze and map three distinct keyword tiers directly to "
         "critical landing pages. This balances immediate traffic through "
         "low-difficulty terms while continuously building long-term "
         "authority on highly competitive keywords."),
        ("Performance Optimization:",
         " Our team runs monthly technical audits and local map listing "
         "optimization. We build domain trust through targeted backlink "
         "acquisition."),
        ("Content & Analytics:",
         " We research and publish two to three targeted blog posts monthly "
         "to capture organic traffic and valuable long-tail queries. Every "
         "month begins with a transparent progress report demonstrating "
         "measurable ranking velocity and traffic growth."),
    ],
    "strategy": {
        "Core SEO":
            "Core SEO sends more qualified traffic and conversions to your "
            "website, improving visibility across search engines (Google) "
            "through on-site optimization, content writing, and ongoing "
            "performance enhancements.",
        "AI Search":
            "AI Search adds on to your traditional search optimization to "
            "Generative Engine Optimization (GEO) to increase your brand’s "
            "visibility and influence within AI-driven search (Google AI "
            "Overviews, ChatGPT, Gemini & more), ensuring your brand is "
            "referenced and recommended.",
        "Website Audit":
            "A comprehensive analysis to uncover technical and content issues "
            "hurting your site’s SEO performance, with an action plan of "
            "what to do.",
    },
    "kw_lead": "Based on our initial research, we came up with a preliminary "
               "list of potential keywords and noted your current rankings "
               "for each keyword on Google.",
    "kw_legend": [
        ("Ultra-Competitive Keywords:",
         " These are the most competitive terms in a particular industry and "
         "are extremely difficult to rank but yield extremely high traffic."),
        ("Competitive Keywords:",
         " These are highly competitive terms which take a high amount of "
         "effort and time to rank for but can drive a tremendous amount of "
         "traffic."),
        ("Long-tail Keywords:",
         " These are lower competition terms which take less time and effort "
         "to rank for but also drive less traffic."),
    ],
    "audit_lead": "Our Comprehensive SEO & Generative Engine Optimization "
                  "Audit provides an in-depth, 25+ page evaluation of your "
                  "website’s organic search performance, technical "
                  "health, content quality, authority and visibility across "
                  "both traditional search engines and emerging AI-powered "
                  "platforms.",
    "audit_bar": "What Is Included",
    "audit": [
        ("Executive Summary and Strategic Recommendations",
         "High-level overview of current search position and key audit findings."),
        ("Analytics, Tracking, and Measurement",
         "Evaluation of measurement systems for search performance, user "
         "behavior, and conversions."),
        ("Technical SEO, Crawling, and Indexation",
         "Review of how efficiently search engines and AI crawlers access and "
         "index the site."),
        ("Website Performance and Core Web Vitals",
         "Assessment of loading speed, usability, and technical ranking factors."),
        ("On-Page SEO and Content Optimization",
         "Evaluation of page structure and relevance for target searches and "
         "user needs."),
        ("Mobile SEO and User Experience",
         "Review of mobile usability, engagement, and conversion factors."),
        ("Structured Data and Schema Markup",
         "Audit of schema implementation for search engines and AI comprehension."),
        ("E-E-A-T and Trust Assessment",
         "Assessment of site experience, expertise, authority, and trustworthiness."),
        ("Generative Engine Optimization and AI-Search Readiness",
         "Evaluation of positioning for discovery and citation by LLMs and AI "
         "search engines."),
        ("Visibility Across Major AI and Search Platforms",
         "Recommendations to enhance presence across major LLMs and search tools."),
        ("Citation-Worthy Content and Topical Authority",
         "Review of content value for external referencing by platforms and "
         "journalists."),
        ("Off-Page SEO, Backlinks, and Authority",
         "Analysis of domain authority, backlink profiles, and competitive gaps."),
        ("Detailed Findings and Supporting Evidence",
         "Structured, evidence-based documentation of audit results."),
        ("Prioritized Next Steps",
         "Actionable roadmap and recommendations for implementation."),
    ],
    "product_title": "Search Engine Optimization+ Product Details",
    # WHAT EACH TIER BUYS. The AI lines are only printed on a quote that
    # carries AI Search -- they are what the AI Search money is for.
    "tiers": [
        {"key": "base", "name": "Base Campaign",
         "ticks": ["In-Depth Monthly Report", "Dedicated SEO Manager",
                   "Increased site traffic"],
         "ai_ticks": ["AI model brand optimization"],
         "last": "Ranking improvements within 6-10 months",
         "counts": [("1", " ultra-competitive keyword"),
                    ("3-5", " competitive keywords"),
                    ("10-15", " long tail keywords")],
         "ai_counts": []},
        {"key": "intermediate", "name": "Intermediate Campaign",
         "ticks": ["In-Depth Monthly Report", "Dedicated SEO Manager",
                   "Increased site traffic"],
         "ai_ticks": ["AI model brand optimization"],
         "then": ["Additional types of link building", "Proprietary rank signaling"],
         "last": "Ranking improvements within 5-8 months",
         "counts": [("2", " ultra-competitive keywords"),
                    ("5-8", " competitive keywords"),
                    ("12-18", " long tail keywords")],
         "ai_counts": [("1-2", " premium placements per month within AI "
                               "search results")]},
        {"key": "advanced", "name": "Advanced Campaign",
         "ticks": ["In-Depth Monthly Report", "Dedicated SEO Manager",
                   "Increased site traffic"],
         "ai_ticks": ["AI model brand optimization"],
         "then": ["Additional types of link building", "Proprietary rank signaling",
                  "Accelerated rank growth and lead generation"],
         "last": "Ranking improvements within 4-6 months",
         "counts": [("3", " ultra-competitive keywords"),
                    ("8-12", " competitive keywords"),
                    ("15-20", " long tail keywords")],
         "ai_counts": [("3-4", " premium placements per month within AI "
                               "search results")]},
    ],
    "term": "Each tier requires a {n}-month term which then becomes a "
            "month-to-month commitment.",
    "no_serp": "No SERP captured for this quote.",
}

TIER_ORDER = ("base", "intermediate", "advanced")
LOGO = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                    "static", "adtini-logo.png")


# ---------------------------------------------------------------- helpers
def _money(v):
    try:
        return "${:,.0f}".format(float(v))
    except (TypeError, ValueError):
        return ""


def _txt(shape, size=12, bold=False, color=INK, align=PP_ALIGN.LEFT,
         space_after=4, line=1.15):
    """One text frame, adtini's type, with the placeholder paragraph reused."""
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.12)
    tf.margin_top = tf.margin_bottom = Inches(0.06)
    p = tf.paragraphs[0]
    p.alignment = align
    p.space_after = Pt(space_after)
    p.line_spacing = line
    _style(p, size, bold, color)
    return tf


def _style(p, size, bold, color):
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.font.name = FONT


def _run(p, text, size=12, bold=False, color=INK):
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    r.font.name = FONT
    return r


def _para(tf, size=12, bold=False, color=INK, space_before=0, space_after=4,
          line=1.15, align=PP_ALIGN.LEFT):
    p = tf.add_paragraph()
    p.alignment = align
    p.space_before = Pt(space_before)
    p.space_after = Pt(space_after)
    p.line_spacing = line
    _style(p, size, bold, color)
    return p


def _box(slide, x, y, w, h, fill=None, outline=LINE, radius=0.04,
         shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    sh = slide.shapes.add_shape(shape, x, y, w, h)
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
    if outline is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = outline
        sh.line.width = Pt(1)
    sh.shadow.inherit = False
    try:
        sh.adjustments[0] = radius
    except (IndexError, KeyError, ValueError):
        pass
    sh.text_frame.word_wrap = True
    sh.text_frame.vertical_anchor = MSO_ANCHOR.TOP
    return sh


def _text_box(slide, x, y, w, h):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tb.text_frame.word_wrap = True
    tb.shadow.inherit = False
    return tb


def _blank(prs):
    """A slide with the adtini chrome: white, navy foot, logo top right."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H)
    bg.fill.solid()
    bg.fill.fore_color.rgb = WHITE
    bg.line.fill.background()
    bg.shadow.inherit = False
    foot = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(7.02),
                                  SLIDE_W, Inches(0.48))
    foot.fill.solid()
    foot.fill.fore_color.rgb = NAVY
    foot.line.fill.background()
    foot.shadow.inherit = False
    # THE MARK, IF THE MARK IS ON DISK. The wordmark is a brand asset rather
    # than something to redraw in shapes; without it the deck prints the name
    # instead of a broken picture, and drops the file in the moment it exists.
    drawn = False
    if os.path.exists(LOGO):
        try:
            slide.shapes.add_picture(LOGO, Inches(11.55), Inches(0.3),
                                     width=Inches(1.35))
            drawn = True
        except Exception:                                     # noqa: BLE001
            drawn = False
    if not drawn:
        tb = _text_box(slide, Inches(11.1), Inches(0.26), Inches(1.9),
                       Inches(0.5))
        _txt(tb, size=20, bold=True, color=NAVY, align=PP_ALIGN.RIGHT)
        tb.text_frame.paragraphs[0].text = "adtini"
    return slide


def _heading(slide, text, size=26, top=0.45):
    tb = _text_box(slide, Inches(0.55), Inches(top), Inches(10.4), Inches(0.75))
    _txt(tb, size=size, bold=True, color=NAVY)
    tb.text_frame.paragraphs[0].text = text
    return tb


def _label_para(tf, label, body, size=11.5, first=False, space_before=6,
                label_color=INK):
    p = tf.paragraphs[0] if first else _para(tf, size=size,
                                             space_before=space_before)
    if first:
        p.space_before = Pt(0)
        p.space_after = Pt(4)
        p.line_spacing = 1.15
    _run(p, label, size=size, bold=True, color=label_color)
    _run(p, body, size=size)
    return p


# ---------------------------------------------------------------- slide 1
def _slide_product(prs, d):
    slide = _blank(prs)
    icon = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.5), Inches(0.42),
                                  Inches(0.72), Inches(0.72))
    icon.fill.solid()
    icon.fill.fore_color.rgb = RGBColor(0xE8, 0xF0, 0xFB)
    icon.line.color.rgb = NAVY
    icon.line.width = Pt(2.5)
    icon.shadow.inherit = False
    _txt(icon, size=20, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
    icon.text_frame.paragraphs[0].text = "⌕"

    tb = _text_box(slide, Inches(1.3), Inches(0.4), Inches(9.6), Inches(0.9))
    _txt(tb, size=32, bold=True, color=NAVY)
    tb.text_frame.paragraphs[0].text = TITLE

    tag = _box(slide, Inches(0.5), Inches(1.3), Inches(6.4), Inches(0.85))
    _txt(tag, size=12.5, color=INK)
    tag.text_frame.paragraphs[0].text = COPY["tagline"]

    # THE EXHIBIT. The capture is the client's own search result, which is the
    # reason the first slide exists; a quote without one says so rather than
    # printing an empty frame.
    img = (d.get("serp") or {}).get("bytes")
    frame = _box(slide, Inches(0.5), Inches(2.4), Inches(6.4), Inches(3.9),
                 fill=RGBColor(0xF4, 0xF5, 0xF7))
    if img:
        pic = slide.shapes.add_picture(io.BytesIO(img), Inches(0.62),
                                       Inches(2.55), width=Inches(6.16))
        # Keep it inside the frame whatever the capture's aspect.
        if pic.height > Inches(3.6):
            ratio = pic.width / pic.height
            pic.height = Inches(3.6)
            pic.width = Emu(int(Inches(3.6) * ratio))
            pic.left = Inches(0.62) + Emu(int((Inches(6.16) - pic.width) / 2))
    else:
        _txt(frame, size=12, color=MUTED, align=PP_ALIGN.CENTER)
        frame.text_frame.paragraphs[0].text = COPY["no_serp"]
        frame.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE

    head = _text_box(slide, Inches(7.2), Inches(1.28), Inches(5.6), Inches(0.4))
    _txt(head, size=15, bold=True, color=INK)
    head.text_frame.paragraphs[0].text = COPY["include_heading"]

    body = _box(slide, Inches(7.2), Inches(1.75), Inches(5.6), Inches(4.55))
    tf = _txt(body, size=11.5)
    for i, (label, text) in enumerate(COPY["include"]):
        _label_para(tf, label, text, size=11.5, first=(i == 0), space_before=9)
    return slide


# ---------------------------------------------------------------- slide 2
def _slide_strategy(prs, strategies):
    slide = _blank(prs)
    _heading(slide, TITLE + " - Strategy Details")
    rows = [(s, COPY["strategy"][s]) for s in ("Core SEO", "AI Search",
                                               "Website Audit")
            if s in strategies]
    y = Inches(1.6)
    for name, text in rows:
        h = Inches(1.0) if len(text) > 200 else Inches(0.8)
        box = _box(slide, Inches(0.55), y, Inches(12.2), h)
        tf = _txt(box, size=12.5)
        _label_para(tf, name + ":", " " + text, size=12.5, first=True,
                    label_color=BLUE)
        y = y + h + Inches(0.22)
    return slide


# ---------------------------------------------------------------- slide 3
def _tier_of(term, kw):
    """Which tier a term was built into -- the list it came out of."""
    t = str(term or "").strip().lower()
    for key, label in (("ultra", "Ultra-Competitive"),
                       ("competitive", "Competitive"),
                       ("long_tail", "Long-Tail")):
        for x in (kw.get(key) or []):
            if str((x or {}).get("kw", "")).strip().lower() == t:
                return label
    return "Competitive"


def _rank_of(term, ranks):
    v = ranks.get(str(term or "").strip().lower())
    if v in (None, "", "—"):
        return "Not Found"
    return str(v)


def _slide_keywords(prs, d, kw, per_slide=12):
    """The list and its split. Paginated, because the list is the engagement:
    a table cut to fit one slide quotes a smaller campaign than the one sold.
    """
    terms = [x for x in (kw.get("all") or []) if (x or {}).get("kw")]
    if not terms:
        return []
    ranks = {}
    for row in (d.get("table") or []):
        k = str((row or {}).get("kw", "")).strip().lower()
        if not k:
            continue
        pos = row.get("pos")
        ranks[k] = pos if pos not in (None, "", 0) else "Not Found"

    counts = {"Ultra-Competitive": len(kw.get("ultra") or []),
              "Competitive": len(kw.get("competitive") or []),
              "Long-Tail": len(kw.get("long_tail") or [])}
    made = []
    pages = [terms[i:i + per_slide] for i in range(0, len(terms), per_slide)]
    for n, page in enumerate(pages):
        slide = _blank(prs)
        _heading(slide, TITLE + " - Keyword Details")
        lead = _text_box(slide, Inches(0.55), Inches(1.24), Inches(12.2),
                         Inches(0.4))
        _txt(lead, size=11.5, color=INK)
        lead.text_frame.paragraphs[0].text = COPY["kw_lead"]

        rows = len(page) + 1
        table = slide.shapes.add_table(
            rows, 3, Inches(0.55), Inches(1.75), Inches(6.9),
            Inches(0.32) * rows).table
        table.columns[0].width = Inches(3.7)
        table.columns[1].width = Inches(1.4)
        table.columns[2].width = Inches(1.8)
        for c, label in enumerate(("Keyword", "Google Rank", "Competitiveness")):
            cell = table.cell(0, c)
            cell.text = label
            cell.fill.solid()
            cell.fill.fore_color.rgb = NAVY
            _style(cell.text_frame.paragraphs[0], 11, True, WHITE)
        for r, x in enumerate(page, start=1):
            term = x.get("kw")
            tier = _tier_of(term, kw)
            for c, val in enumerate((term, _rank_of(term, ranks), tier)):
                cell = table.cell(r, c)
                cell.text = str(val)
                cell.fill.solid()
                cell.fill.fore_color.rgb = WHITE
                p = cell.text_frame.paragraphs[0]
                bold = c == 2 and tier != "Competitive"
                color = BLUE_LT if (c == 2 and tier == "Competitive") else INK
                if c == 2 and tier != "Competitive":
                    color = NAVY
                _style(p, 10.5, bold, color)
                p.alignment = PP_ALIGN.CENTER if c == 1 else PP_ALIGN.LEFT

        # The split, once -- it describes the whole list, not the page.
        if n == 0:
            _distribution(slide, counts)
        made.append(slide)
    return made


def _distribution(slide, counts):
    head = _text_box(slide, Inches(7.7), Inches(1.68), Inches(5.1), Inches(0.4))
    _txt(head, size=15, bold=True, color=INK)
    head.text_frame.paragraphs[0].text = "Keyword Distribution"

    data = CategoryChartData()
    labels = [k for k in ("Ultra-Competitive", "Competitive", "Long-Tail")
              if counts.get(k)]
    data.categories = labels
    data.add_series("Keywords", tuple(counts[k] for k in labels))
    gf = slide.shapes.add_chart(XL_CHART_TYPE.PIE, Inches(7.55), Inches(2.05),
                                Inches(5.3), Inches(2.5), data)
    chart = gf.chart
    chart.has_title = False
    chart.has_legend = True
    chart.legend.include_in_layout = False
    chart.legend.font.size = Pt(11)
    chart.legend.font.name = FONT
    plot = chart.plots[0]
    plot.has_data_labels = True
    plot.data_labels.show_value = True
    plot.data_labels.font.size = Pt(10)
    plot.data_labels.font.bold = True
    plot.data_labels.font.color.rgb = WHITE
    plot.data_labels.font.name = FONT
    plot.data_labels.position = XL_LABEL_POSITION.INSIDE_END
    shades = {"Ultra-Competitive": NAVY, "Competitive": BLUE_LT,
              "Long-Tail": BLUE_MID}
    for i, label in enumerate(labels):
        pt = plot.series[0].points[i]
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = shades[label]
        pt.format.line.color.rgb = WHITE

    legend = _text_box(slide, Inches(7.55), Inches(4.6), Inches(5.3),
                       Inches(2.2))
    tf = _txt(legend, size=10.5)
    for i, (label, text) in enumerate(COPY["kw_legend"]):
        _label_para(tf, "•  " + label, text, size=10.5, first=(i == 0),
                    space_before=7)


# ---------------------------------------------------------------- slide 4
def _slide_audit(prs):
    slide = _blank(prs)
    _heading(slide, TITLE + " - Audit Details")
    lead = _text_box(slide, Inches(0.55), Inches(1.22), Inches(12.2),
                     Inches(0.6))
    _txt(lead, size=11.5, color=INK)
    lead.text_frame.paragraphs[0].text = COPY["audit_lead"]

    bar = _box(slide, Inches(0.55), Inches(1.95), Inches(12.2), Inches(0.42),
               fill=NAVY, outline=None)
    _txt(bar, size=12.5, bold=True, color=WHITE)
    bar.text_frame.paragraphs[0].text = COPY["audit_bar"]
    bar.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE

    body = _box(slide, Inches(0.55), Inches(2.37), Inches(12.2), Inches(4.35))
    tf = _txt(body, size=10.5)
    for i, (label, text) in enumerate(COPY["audit"]):
        _label_para(tf, "•  " + label + ":", " " + text, size=10.5,
                    first=(i == 0), space_before=3.5, label_color=BLUE)
    return slide


# ---------------------------------------------------------------- slide 5
# HOW TALL A ROW IS DEPENDS ON THE WORDS IN IT. Every row was given the same
# fixed height, and "In-Depth Monthly Report" wraps to two lines in the card's
# left column -- so each tick sat on top of the one below it and the last of
# them ran off the bottom of the card. (2026-09-19, Kiri)
def _lines(text, per_line):
    line, n = 0, 1
    for word in str(text).split():
        add = len(word) + (1 if line else 0)
        if line + add > per_line and line:
            n += 1
            line = len(word)
        else:
            line += add
    return n


def _tick(slide, x, y, w, text, tint=None, bold=False, color=None):
    h = Inches(0.155) * _lines(text, 26) + Inches(0.05)
    tb = _text_box(slide, x, y, w - Inches(0.3), h)
    _txt(tb, size=8.5, color=(color or tint or INK), space_after=0, line=1.0)
    tb.text_frame.paragraphs[0].text = text
    tb.text_frame.paragraphs[0].font.bold = bold
    mark = _text_box(slide, x + w - Inches(0.32), y, Inches(0.28), Inches(0.24))
    _txt(mark, size=10, bold=True, color=GREEN, align=PP_ALIGN.CENTER,
         space_after=0)
    mark.text_frame.paragraphs[0].text = "✓"
    return h + Inches(0.07)


def _priced(d):
    """Whether there is a price to put on a pricing slide at all."""
    pricing = d.get("pricing") or {}
    h = pricing.get("handoff") or {}
    tiers = (h.get("core_seo_price") or pricing.get("client_tiers") or {})
    ai = (h.get("ai_search_price")
          or (pricing.get("ai_search") or {}).get("client_add") or {})
    return any(tiers.get(k) or ai.get(k) for k in TIER_ORDER)


def _slide_pricing(prs, d):
    pricing = d.get("pricing") or {}
    h = pricing.get("handoff") or {}
    core = h.get("core_seo_price") or pricing.get("client_tiers") or {}
    ai = h.get("ai_search_price") or (
        (pricing.get("ai_search") or {}).get("client_add")) or {}
    ai_on = any(ai.get(k) for k in TIER_ORDER)
    # A QUOTE WITH AI SEARCH AND NO CORE SEO PRICES THE AI LEG AS THE CAMPAIGN.
    # Then the headline is that number and there is no second line under it.
    core_on = any(core.get(k) for k in TIER_ORDER)
    if not core_on and ai_on:
        core, ai, ai_on = ai, {}, False
    term = int(pricing.get("min_term_months") or h.get("min_term_months") or 6)
    n_addon = int(pricing.get("addon_markets") or h.get("addon_markets") or 0)
    per_market = (pricing.get("client_addon_per_market")
                  or h.get("addon_market_price") or {})

    slide = _blank(prs)
    icon = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.5), Inches(0.44),
                                  Inches(0.7), Inches(0.7))
    icon.fill.solid()
    icon.fill.fore_color.rgb = RGBColor(0xE8, 0xF0, 0xFB)
    icon.line.color.rgb = NAVY
    icon.line.width = Pt(2.5)
    icon.shadow.inherit = False
    _txt(icon, size=19, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
    icon.text_frame.paragraphs[0].text = "⌕"
    tb = _text_box(slide, Inches(1.26), Inches(0.44), Inches(10.4), Inches(0.8))
    _txt(tb, size=27, bold=True, color=NAVY)
    tb.text_frame.paragraphs[0].text = COPY["product_title"]

    # ---- the summary box. Each cell is printed only where the quote has one.
    monthly = (core.get("intermediate") or 0) + (ai.get("intermediate") or 0)
    cells = [[("Months Running: ", str(term))]]
    if monthly:
        cells.append([("Monthly Budget: ", _money(monthly)),
                      ("Total Budget: ", _money(monthly * term))])
    if n_addon and per_market.get("intermediate"):
        cells.append([("# of Add-on Markets: ", str(n_addon)),
                      ("Add-on Price: ", _money(per_market["intermediate"]))])
    top = _box(slide, Inches(0.55), Inches(1.32), Inches(12.2), Inches(0.72))
    span = Inches(12.2) / len(cells)
    for i, lines in enumerate(cells):
        cell = _text_box(slide, Inches(0.55) + span * i, Inches(1.38), span,
                         Inches(0.62))
        tf = _txt(cell, size=12, align=PP_ALIGN.CENTER, space_after=0)
        for j, (label, value) in enumerate(lines):
            p = tf.paragraphs[0] if j == 0 else _para(
                tf, size=12, align=PP_ALIGN.CENTER, space_after=0)
            p.alignment = PP_ALIGN.CENTER
            _run(p, label, size=12, bold=True)
            _run(p, value, size=12)
        if i:
            rule = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE, Inches(0.55) + span * i, Inches(1.4),
                Pt(1), Inches(0.56))
            rule.fill.solid()
            rule.fill.fore_color.rgb = LINE
            rule.line.fill.background()
            rule.shadow.inherit = False

    # ---- the three campaign cards
    card_w = Inches(3.95)
    gap = Inches(0.22)
    left0 = Inches(0.55)
    for i, spec in enumerate(COPY["tiers"]):
        key = spec["key"]
        if not core.get(key):
            continue
        x = left0 + (card_w + gap) * i
        dark = i > 0
        fill = None if i == 0 else (BLUE_MID if i == 1 else BLUE_LT)
        card = _box(slide, x, Inches(2.12), card_w, Inches(4.55),
                    fill=fill, outline=(LINE if i == 0 else None))
        name = _text_box(slide, x, Inches(2.22), card_w, Inches(0.42))
        _txt(name, size=15, bold=True,
             color=(INK if not dark else WHITE), align=PP_ALIGN.CENTER)
        name.text_frame.paragraphs[0].text = spec["name"]

        price = _text_box(slide, x, Inches(2.66), card_w, Inches(0.42))
        tf = _txt(price, size=20, align=PP_ALIGN.CENTER, space_after=0)
        _run(tf.paragraphs[0], _money(core[key]), size=20, bold=True,
             color=(BLUE if not dark else WHITE))
        _run(tf.paragraphs[0], " / month", size=12,
             color=(INK if not dark else WHITE))
        head_h = Inches(1.02)
        if ai_on and ai.get(key):
            sub = _text_box(slide, x, Inches(3.08), card_w, Inches(0.3))
            tf = _txt(sub, size=11, align=PP_ALIGN.CENTER, space_after=0)
            _run(tf.paragraphs[0], "AI Search: ", size=11, bold=True,
                 color=(INK if not dark else WHITE))
            _run(tf.paragraphs[0], _money(ai[key]) + " / month", size=11,
                 color=(INK if not dark else WHITE))
            head_h = Inches(1.34)

        rule = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x,
                                      Inches(2.12) + head_h, card_w, Pt(1))
        rule.fill.solid()
        rule.fill.fore_color.rgb = LINE if not dark else WHITE
        rule.line.fill.background()
        rule.shadow.inherit = False
        mid = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x + Inches(2.05),
                                     Inches(2.12) + head_h + Inches(0.1),
                                     Pt(1), Inches(4.55) - head_h - Inches(0.2))
        mid.fill.solid()
        mid.fill.fore_color.rgb = LINE if not dark else WHITE
        mid.line.fill.background()
        mid.shadow.inherit = False

        ticks = list(spec["ticks"])
        if ai_on:
            ticks += spec["ai_ticks"]
        ticks += list(spec.get("then") or [])
        tint = TINT if dark else MUTED
        y = Inches(2.12) + head_h + Inches(0.13)
        for t in ticks:
            y = y + _tick(slide, x + Inches(0.1), y, Inches(1.95), t, tint=tint)
        _tick(slide, x + Inches(0.1), y, Inches(1.95), spec["last"],
              bold=True, color=(BLUE if not dark else WHITE))

        counts = list(spec["counts"])
        if ai_on:
            counts += spec["ai_counts"]
        cy = Inches(2.12) + head_h + Inches(0.13)
        for n, rest in counts:
            ch = Inches(0.155) * _lines(n + rest, 24) + Inches(0.05)
            cb = _text_box(slide, x + Inches(2.16), cy, Inches(1.68), ch)
            tf = _txt(cb, size=8.5, space_after=0, line=1.0)
            _run(tf.paragraphs[0], n, size=8.5, bold=True,
                 color=(INK if not dark else WHITE))
            _run(tf.paragraphs[0], rest, size=8.5,
                 color=(INK if not dark else WHITE))
            cy = cy + ch + Inches(0.07)

    foot = _text_box(slide, Inches(0.55), Inches(6.74), Inches(9.4), Inches(0.32))
    _txt(foot, size=11, bold=True, color=INK)
    foot.text_frame.paragraphs[0].text = COPY["term"].format(n=term)
    return slide


# ---------------------------------------------------------------- build
def build_proposal_pptx(d):
    """The deck, from the same quote the proposal document is built from."""
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    strategies = _strategies(d)
    kw = d.get("kw") or {}

    _slide_product(prs, d)
    if strategies:
        _slide_strategy(prs, strategies)
    # THE KEYWORD SLIDE BELONGS TO THE SEARCH STRATEGIES. An audit-only
    # engagement buys a report, not a keyword set, and the tier split would be
    # describing a campaign nobody bought.
    if strategies & {"Core SEO", "AI Search"}:
        _slide_keywords(prs, d, kw)
    if "Website Audit" in strategies:
        _slide_audit(prs)
    # A PRICING SLIDE WITH NO PRICE ON IT IS A SLIDE TO DELETE BY HAND.
    if _priced(d):
        _slide_pricing(prs, d)

    buf = io.BytesIO()
    prs.save(buf)
    buf.seek(0)
    return buf


def _strategies(d):
    """What was sold, however the quote spells it."""
    raw = d.get("strategy")
    if isinstance(raw, str):
        parts = [x.strip() for x in raw.split(",")]
    else:
        parts = [str(x).strip() for x in (raw or [])]
    known = {"core seo": "Core SEO", "ai search": "AI Search",
             "website audit": "Website Audit"}
    return {known[p.lower()] for p in parts if p.lower() in known}
