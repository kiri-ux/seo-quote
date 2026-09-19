"""THE DECK IS THE QUOTE, SLIDE BY SLIDE.

The proposal document is the letter; the adtini product slides are what the
partner presents, and they were being rebuilt by hand from the product deck
every time -- which is where a Website Audit section reaches a client who
never bought one, and an AI Search price reaches one who is only buying Core
SEO.

So every slide here is checked against what the quote carries: the strategies
that were sold and no others, the keyword slide only for the search
strategies, the audit slide only for the audit, and on the pricing slide the
AI Search line and the add-on market box only when the quote holds them.
(2026-09-19, Kiri)
"""
import importlib.util
import io
import os
import sys

os.environ.setdefault("DFS_LOGIN", "x")
os.environ.setdefault("DFS_PASSWORD", "x")
HERE = os.path.dirname(os.path.abspath(__file__))
SRCDIR = os.path.dirname(HERE)
sys.path.insert(0, SRCDIR)
SRC = os.path.join(SRCDIR, "app.py")
spec = importlib.util.spec_from_file_location("app", SRC)
app = importlib.util.module_from_spec(spec)
sys.modules["app"] = app
spec.loader.exec_module(app)
import seo_pptx as P      # noqa: E402
from pptx import Presentation      # noqa: E402

FAIL, CHECKS = [], []
TITLE_LINE = "Search Engine Optimization+"


def check(label, got, want):
    ok = got == want
    CHECKS.append(label)
    print(("  ok   " if ok else "  FAIL ") + label)
    if not ok:
        print("         got  %r\n         want %r" % (got, want))
        FAIL.append(label)


TERMS = ["apartments for rent san diego ca", "homes for sale san diego ca",
         "luxury apartments san diego ca", "homes for rent san diego ca",
         "new homes san diego ca", "studio apartments san diego ca"]
KW = {"ultra": [{"kw": TERMS[0]}, {"kw": TERMS[1]}],
      "competitive": [{"kw": TERMS[2]}, {"kw": TERMS[3]}],
      "long_tail": [{"kw": TERMS[4]}, {"kw": TERMS[5]}],
      "all": [{"kw": t} for t in TERMS]}


def quote(strategy, ai=True, addon=3, term=12):
    p = app.stage4_price("single_city", 550, False, addon, 35,
                         pct_not_ranking=100, total_volume=9220,
                         ai_search=ai, core_seo=True)
    return {"brand": "Escaya", "order_no": "44973", "strategy": strategy,
            "kw": KW,
            "table": [{"kw": TERMS[0], "pos": 3}, {"kw": TERMS[1], "pos": None}],
            "pricing": dict(p, addon_markets=addon, min_term_months=term,
                            client_addon_per_market=p.get(
                                "client_addon_per_market")),
            "serp": {"bytes": None}}


def deck(d):
    return Presentation(P.build_proposal_pptx(d))


def slides(prs):
    """Each slide as one flat string, which is how a reader meets it."""
    out = []
    for s in prs.slides:
        parts = []
        for sh in s.shapes:
            if sh.has_text_frame and sh.text_frame.text.strip():
                parts.append(sh.text_frame.text)
            if getattr(sh, "has_table", False) and sh.has_table:
                for row in sh.table.rows:
                    parts += [c.text for c in row.cells]
        out.append(" · ".join(parts))
    return out


print("\nEVERYTHING SOLD, AND NOTHING ELSE")
full = slides(deck(quote("Core SEO, AI Search, Website Audit")))
check("five slides on the full quote", len(full), 5)
check("the product and its SERP first",
      TITLE_LINE in full[0] and "Your SEO Campaign Will Include:" in full[0],
      True)
check("a capture that was never taken says so",
      "No SERP captured" in full[0], True)
check("strategy details second", "Strategy Details" in full[1], True)
check("with all three strategies on it",
      all(s + ":" in full[1] for s in
          ("Core SEO", "AI Search", "Website Audit")), True)
check("keyword details third", "Keyword Details" in full[2], True)
check("audit details fourth", "Audit Details" in full[3], True)
check("and the product details last", "Product Details" in full[4], True)

print("\nA STRATEGY THAT WAS NOT SOLD IS NOT IN THE DECK")
core = slides(deck(quote("Core SEO", ai=False, addon=0, term=6)))
check("no audit slide", any("Audit Details" in s for s in core), False)
check("and no audit line on the strategy slide",
      "Website Audit:" in core[1], False)
check("nor an AI Search line", "AI Search:" in core[1], False)
check("the keyword slide stays, because SEO was sold",
      any("Keyword Details" in s for s in core), True)

audit = slides(deck(quote("Website Audit", ai=False, addon=0)))
check("an audit-only quote has no keyword slide",
      any("Keyword Details" in s for s in audit), False)
check("but keeps the audit slide",
      any("Audit Details" in s for s in audit), True)

print("\nTHE AI SEARCH PRICE ONLY WHEN IT IS IN THE QUOTE")
with_ai = slides(deck(quote("Core SEO, AI Search")))[-1]
# THE HEADLINE IS WHAT THEY PAY. It was the Core SEO price with AI Search
# printed beside it, so the Monthly Budget above -- both legs -- matched no
# card on the slide.
check("the tier price is both legs together", "$6,900 / month" in with_ai, True)
check("with Core SEO named under it",
      "Core SEO: $3,950 / month" in with_ai, True)
check("and AI Search under that",
      "AI Search: $2,950 / month" in with_ai, True)
check("and the AI lines in the tier list",
      "AI model brand optimization" in with_ai, True)
check("and what the AI money buys",
      "premium placements per month within AI search results" in with_ai, True)
check("no AI Search line without it", "AI Search:" in core[-1], False)
check("and no split at all, because there is nothing to split",
      "Core SEO: $" in core[-1], False)
check("and none of its bullets either",
      "AI model brand optimization" in core[-1], False)
# A quote that buys AI Search alone prices THAT as the campaign rather than
# printing a second line under a price that is not there.
solo = app.stage4_price("single_city", 550, False, 0, 35, pct_not_ranking=100,
                        total_volume=9220, ai_search=True, core_seo=False)
d = quote("AI Search")
d["pricing"] = dict(solo, addon_markets=0, min_term_months=12)
only_ai = slides(deck(d))[-1]
check("an AI-only quote prices the AI campaign", "/ month" in only_ai, True)
check("and carries no second price line under it",
      "AI Search: $" in only_ai, False)

print("\nADD-ON MARKETS ONLY WHEN THEY ARE ON THE QUOTE")
check("the count and the rate are printed",
      "# of Add-on Markets: 3" in with_ai.replace("\n", " ")
      or "# of Add-on Markets: 3" in slides(deck(quote(
          "Core SEO, AI Search")))[-1], True)
check("and neither is printed without them",
      "Add-on" in core[-1], False)

print("\nTHE NUMBERS ARE THE QUOTE'S OWN")
check("the term is the quote's", "12-month term" in with_ai, True)
check("a six-month quote says six", "6-month term" in core[-1], True)
# Monthly budget is the intermediate tier, which is the headline the tool
# prints everywhere else, times the term it is sold on.
check("the budget is the intermediate tier",
      "Monthly Budget: $8,650" in with_ai, True)
check("and the total is that across the term",
      "Total Budget: $103,800" in with_ai, True)
# A FILE CANNOT BE SELECTED ON SCREEN, so the tier the budget came from is
# marked on its own card.
check("the tier is marked on its own card", "Selected" in with_ai, True)
check("the budget equals that card", "$8,650 / month" in with_ai, True)

print("\nTHE FLIGHT IS WHAT WAS TYPED ON THE FORM")
flight = dict(quote("Core SEO, AI Search"),
              start_date="2025-05-01", end_date="2025-10-31", months="")
dated = slides(deck(flight))[-1]
check("the dates print across the top",
      "05/01/2025 - 10/31/2025" in dated, True)
check("and the months are read off them", "Months Running: 6" in dated, True)
check("the total is the budget across those months",
      "Total Budget: $51,900" in dated, True)
# Typed months win, and an end date is worked out from them.
typed = dict(quote("Core SEO, AI Search"), start_date="2025-05-01", months="3")
check("months typed by hand set the end date",
      "05/01/2025 - 08/01/2025" in slides(deck(typed))[-1], True)
check("with no dates at all it falls back to the term",
      "Months Running: 12" in with_ai, True)
check("and never prints a range it does not have",
      "/2025 -" in with_ai, False)

print("\nTHE AI LINES CARRY THE SPARKLE")
check("the AI tier line is marked with it",
      "\u2726" in with_ai, True)
check("and a quote without AI Search has none",
      "\u2726" in core[-1], False)

print("\nTHE KEYWORD SLIDE IS THE LIST, NOT A SAMPLE OF IT")
kw_slide = slides(deck(quote("Core SEO")))[2]
for t in TERMS:
    check("carries " + t, t in kw_slide, True)
check("with the rank the check found", " 3 " in " " + kw_slide + " ", True)
check("and says so when it found none", "Not Found" in kw_slide, True)
check("the tiers are named",
      all(x in kw_slide for x in ("Ultra-Competitive", "Competitive",
                                  "Long-Tail")), True)
# LONGER THAN A SLIDE IS MORE SLIDES. A table cut to fit quotes a smaller
# engagement than the one being sold.
many = quote("Core SEO")
many["kw"] = dict(KW, all=[{"kw": "term %d" % i} for i in range(30)])
check("thirty terms run onto a second slide",
      sum("Keyword Details" in s for s in slides(deck(many))), 3)

print("\nA CHIP CAN CARRY TWO STRATEGIES")
# The form offers the pair as one chip, "Core SEO + AI Search", and reading
# the field on commas alone matched neither: a real quote came back with the
# product slide and the prices and nothing in between. (2026-09-19, Kiri)
pair = slides(deck(quote("Core SEO + AI Search")))
check("both are read off one chip",
      sorted(P._strategies({"strategy": "Core SEO + AI Search"})),
      ["AI Search", "Core SEO"])
check("a list of chips reads the same",
      sorted(P._strategies({"strategy": ["Core SEO + AI Search"]})),
      ["AI Search", "Core SEO"])
check("so the strategy slide is there",
      any("Strategy Details" in s for s in pair), True)
check("and the keyword slide with it",
      any("Keyword Details" in s for s in pair), True)
check("with both strategies on the strategy slide",
      "Core SEO:" in pair[1] and "AI Search:" in pair[1], True)

print("\nTHE LOGO IS A FILE, NOT A WORD")
check("the deck never sets the name in type",
      any("adtini" in s for s in pair), False)

print("\nAND IT COMES BACK AS A FILE")
app.app.config["TESTING"] = True
client = app.app.test_client()
r = client.post("/api/proposal.pptx", json=quote("Core SEO, AI Search"))
check("the endpoint answers", r.status_code, 200)
check("with a deck", r.data[:2], b"PK")
check("named for the client",
      "Escaya_44973.pptx" in r.headers.get("Content-Disposition", ""), True)
check("and it opens", len(Presentation(io.BytesIO(r.data)).slides) > 0, True)

print("\n%d checks, %d failed" % (len(CHECKS), len(FAIL)))
if FAIL:
    print("FAILED: " + ", ".join(FAIL))
    sys.exit(1)
print("all ok")
