"""A LOOKALIKE DOMAIN IS A FINDING, NOT NOISE.

Wheeler Digital ranks #1 for its own name and wheeler-digital.com ranks #2.
The scan searched only "{brand} reviews", set the lookalike aside as another
company's, and the snapshot said nothing on page one named the client.
Brendan wrote three paragraphs about it by hand. (2026-10-25, Kiri)
"""
import os
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.dirname(HERE))
import rep_scan

FAIL, RUN = [], []


def check(label, got, want):
    ok = got == want
    RUN.append(label)
    print(("  ok   " if ok else "  FAIL ") + label)
    if not ok:
        print("         got  %r\n         want %r" % (got, want))
        FAIL.append(label)


def org(pos, domain, title):
    return {"type": "organic", "rank_group": pos, "domain": domain,
            "url": "https://%s/" % domain, "title": title, "description": ""}


PAGES = {
    "wheeler digital reviews": [
        org(1, "wheeler-digital.com", "Wheeler Digital Solutions Reviews"),
        org(2, "wheelertools.com", "Wheeler Digital Fat Wrench Reviews"),
    ],
    "wheeler digital": [
        org(1, "wheelerdigital.com", "Wheeler Digital | Marketing Agency"),
        org(2, "wheeler-digital.com", "Wheeler Digital Solutions"),
        org(3, "wheelertools.com", "Wheeler Digital Torque Wrench"),
    ],
}
ASKED = []


def fake_post(path, tasks, timeout=None):
    kw = tasks[0]["keyword"]
    ASKED.append(kw)
    return {"tasks": [{"result": [{"items": PAGES.get(kw, [])}]}]}


rep_scan.init(fake_post)
r = rep_scan.scan_serp("Wheeler Digital", "wheelerdigital.com")
check("the bare name is searched as well as the reviews page",
      ASKED, ["wheeler digital reviews", "wheeler digital"])
check("their own rank for the name is read", r["brand_own_rank"], 1)
check("the lookalike is kept, at its best rank on the name search",
      [(x["domain"], x["pos"], x["query"]) for x in r["lookalikes"]],
      [("wheeler-digital.com", 2, "wheeler digital")])
check("a same-word company on another domain is not a lookalike",
      any(x["domain"] == "wheelertools.com" for x in r["lookalikes"]), False)

# A typed query that already says reviews still searches the bare name.
ASKED.clear()
rep_scan.scan_serp("Wheeler Digital", "wheelerdigital.com",
                   query="wheeler digital reviews")
check("a typed reviews query still searches the bare name",
      ASKED[-1], "wheeler digital")


def broken(path, tasks, timeout=None):
    if tasks[0]["keyword"] == "wheeler digital":
        raise RuntimeError("down")
    return fake_post(path, tasks, timeout)


rep_scan.init(broken)
r = rep_scan.scan_serp("Wheeler Digital", "wheelerdigital.com")
check("a failed name search never fails the scan",
      (r["brand_own_rank"], [x["domain"] for x in r["lookalikes"]]),
      (None, ["wheeler-digital.com"]))

# ---------------------------------------------- the proposal and the deck
import rep_docx
import rep_pptx
from pptx import Presentation

SNAP = {"query": "wheeler digital reviews", "organic": [], "forums": [],
        "searched": True, "brand_volume": 90,
        "lookalikes": [{"domain": "wheeler-digital.com", "pos": 2}],
        "brand_own_rank": 1, "locations": [{"title": "Wheeler Digital",
                                            "profile_rating": 5,
                                            "profile_reviews": 7, "pos_5": 7}]}
paras = rep_docx.lookalike_paragraphs(SNAP, "Wheeler Digital",
                                      "https://www.wheelerdigital.com/")
check("one lookalike reads as Brendan wrote it",
      paras[0].startswith("wheelerdigital.com currently ranks #1 organically "
                          "for branded searches, but the similarly named "
                          "wheeler-digital.com ranks #2, creating confusion"),
      True)
check("in three paragraphs, singular throughout",
      (len(paras), "competing domain\u2019s visibility" in paras[0],
       "competing site lower" in paras[1]), (3, True, True))
check("no lookalike, no paragraphs",
      rep_docx.lookalike_paragraphs(dict(SNAP, lookalikes=[]), "X"), [])
off = rep_docx.lookalike_paragraphs(dict(SNAP, brand_own_rank=None),
                                    "Wheeler Digital", "wheelerdigital.com")
check("a client off page one gets the other opening",
      "does not appear on page one" in off[0], True)
two = rep_docx.lookalike_paragraphs(
    dict(SNAP, lookalikes=SNAP["lookalikes"] + [{"domain": "wheelerdigital.net",
                                                 "pos": 5}]),
    "Wheeler Digital", "wheelerdigital.com")
check("two lookalikes are listed with their ranks, plural",
      ("wheeler-digital.com (#2) and wheelerdigital.net (#5)" in two[0],
       "competing domains\u2019 visibility" in two[0]), (True, True))

prs = Presentation(rep_pptx.build_rep_proposal_pptx(
    {"brand": "Wheeler Digital", "site": "https://wheelerdigital.com",
     "quote": {"handoff": {"strategy": ["Proactive"]}}, "snapshot": SNAP}))
txt = ""
for sl in prs.slides:
    for sh in sl.shapes:
        if sh.has_text_frame:
            txt += sh.text_frame.text + "\n"
        if getattr(sh, "has_table", False) and sh.has_table:
            for row in sh.table.rows:
                txt += " ".join(c.text for c in row.cells) + "\n"
check("the deck lists their site and the lookalike on page one",
      ("wheelerdigital.com" in txt and "Owned" in txt,
       "Competing name" in txt), (True, True))
check("their own name rank counts as owned on the Proactive card",
      ("1 of 2", "page-one results they own") in rep_pptx._card_stats(
          {"site": "wheelerdigital.com", "snapshot": SNAP}, "proactive"), True)

BOTH = {"brand": "Wheeler Digital", "snapshot": SNAP,
        "quote": {"handoff": {"strategy": ["Reactive", "Proactive"]}}}
REACT = dict(BOTH, quote={"handoff": {"strategy": ["Reactive"]}})
check("on a bundle, Reactive carries nothing",
      rep_pptx._lookalike_blurb(BOTH, "reactive"), "")
check("and Proactive carries both of Brendan's sentences",
      rep_pptx._lookalike_blurb(BOTH, "proactive"),
      "We will optimize Wheeler Digital\u2019s existing web properties and "
      "build out additional branded profiles, pages, and content, working to "
      "move more Wheeler Digital assets above wheeler-digital.com and push the "
      "competing site lower in search results. Throughout the campaign, we "
      "will monitor branded search rankings, track "
      "wheeler-digital.com\u2019s visibility, and adjust our strategy as "
      "results develop.")
check("Reactive alone gets only the push-lower line",
      rep_pptx._lookalike_blurb(REACT, "reactive"),
      "We will work to push wheeler-digital.com lower in search results.")
check("no lookalike, no line",
      rep_pptx._lookalike_blurb({"snapshot": dict(SNAP, lookalikes=[])},
                                "reactive"), "")

print()
print("%d checks, %d failed" % (len(RUN), len(FAIL)))
print("all ok" if not FAIL else "FAILED: " + ", ".join(FAIL))
sys.exit(1 if FAIL else 0)
