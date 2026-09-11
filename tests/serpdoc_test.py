"""THE SCREENSHOT NEVER REACHED THE DOCUMENT.

The reputation tool has captured real SERP screenshots since the scan
shipped -- Proposal Visuals renders one per query with a Download link -- and
every one of them stopped there. The evidence both SSG reputation proposals
open their search section with (Sage listing what ranks for "Sage Dental
Reviews", Visions showing Related Searches marked in red) had to be saved out
and pasted into the Word file by hand.

The captures live in the browser as data URLs. They now travel with the
download and land under a Search Results heading, captioned with the query
that produced them.
"""
import base64
import io
import os
import struct
import sys
import zlib

HERE = os.path.dirname(os.path.abspath(__file__))
SRCDIR = os.path.dirname(HERE)
sys.path.insert(0, SRCDIR)
os.environ.setdefault("DFS_LOGIN", "x")
os.environ.setdefault("DFS_PASSWORD", "x")
import rep_docx as D          # noqa: E402
import rep_pricing as R       # noqa: E402
from docx import Document     # noqa: E402

FAIL, CHECKS = [], []


def check(label, got, want):
    ok = got == want
    CHECKS.append(label)
    print(("  ok   " if ok else "  FAIL ") + label)
    if not ok:
        print("         got  %r\n         want %r" % (got, want))
        FAIL.append(label)


def png(w, h):
    raw = b"".join(b"\x00" + bytes((200, 210, 230)) * w for _ in range(h))

    def chunk(t, d):
        c = t + d
        return struct.pack(">I", len(d)) + c + struct.pack(">I", zlib.crc32(c) & 0xFFFFFFFF)
    return (b"\x89PNG\r\n\x1a\n"
            + chunk(b"IHDR", struct.pack(">IIBBBBB", w, h, 8, 2, 0, 0, 0))
            + chunk(b"IDAT", zlib.compress(raw))
            + chunk(b"IEND", b""))


IMG = png(600, 400)
URL = "data:image/jpeg;base64," + base64.b64encode(IMG).decode()
QUOTE = R.build_rep_quote({"campaign": "bundle", "margin_pct": 0.35,
                           "search": {"bundle": True, "volume": 8000},
                           "shield": {"locations": 4},
                           "reviews": {"count": 26},
                           "articles": {"standard": 2, "premium": 1}})
BRACKETS = [{"min": 1, "max": 25, "price": 900}]


def build(shots):
    return Document(D.build_rep_proposal_docx(
        {"brand": "Junk Bee Gone", "quote": QUOTE,
         "serp_shots": shots, "brackets": BRACKETS}))


print("\nTHE CAPTURE LANDS IN THE PROPOSAL")
doc = build([{"query": "junk bee gone llc reviews", "data_url": URL}])
texts = [p.text for p in doc.paragraphs if p.text.strip()]
check("one image", len(doc.inline_shapes), 1)
check("sized to the page", round(doc.inline_shapes[0].width.inches, 1), 6.2)
check("under its own heading", "Search Results" in texts, True)
check("captioned with the query",
      "“junk bee gone llc reviews”" in texts, True)

print("\nSEVERAL CAPTURES, CAPPED")
many = [{"query": "q%d" % i, "data_url": URL} for i in range(6)]
check("three is the cap", len(D.serp_images(many)), 3)
check("and the document agrees", len(build(many).inline_shapes), 3)
check("cap is liftable", len(D.serp_images(many, cap=6)), 6)

print("\nNO CAPTURE IS NOT AN EMPTY HEADING")
plain = build(None)
check("no image", len(plain.inline_shapes), 0)
check("no heading either",
      "Search Results" in [p.text for p in plain.paragraphs], False)
check("the proposal still builds",
      any("Summary" == p.text for p in plain.paragraphs), True)

print("\nA BAD CAPTURE IS SKIPPED, NOT FATAL")
# Anything that reaches python-docx malformed raises inside the download.
for label, shots in [
        ("not a data url", [{"query": "x", "data_url": "https://example.com/a.jpg"}]),
        ("undecodable", [{"query": "x", "data_url": "data:image/png;base64,%%%"}]),
        ("not an image", [{"query": "x", "data_url": "data:text/plain;base64,aGk="}]),
        ("no url", [{"query": "x"}]),
        ("not a dict", ["nope"]),
        ("empty", []),
        ("none", None)]:
    check(label, len(D.serp_images(shots)), 0)

print("\nSIZE IS BOUNDED AT BOTH ENDS")
check("a truncated file is refused",
      len(D.serp_images([{"query": "x",
                          "data_url": "data:image/jpeg;base64," +
                          base64.b64encode(b"\xff\xd8" * 8).decode()}])), 0)
_big = "data:image/jpeg;base64," + base64.b64encode(b"\x00" * (13 * 1024 * 1024)).decode()
check("so is an oversized one", len(D.serp_images([{"query": "x", "data_url": _big}])), 0)

print("\nA GOOD ONE BESIDE A BAD ONE STILL GOES")
mixed = build([{"query": "bad", "data_url": "nope"},
               {"query": "good", "data_url": URL}])
check("the good one lands", len(mixed.inline_shapes), 1)
check("only its caption",
      [p.text for p in mixed.paragraphs if p.text.strip() in ("“bad”", "“good”")],
      ["“good”"])

print("\nAN UNNAMED CAPTURE IS STILL A CAPTURE")
anon = build([{"data_url": URL}])
check("image", len(anon.inline_shapes), 1)
check("no empty quotes",
      [p.text for p in anon.paragraphs if p.text.strip() == "“”"], [])

print("\nTHE REMOVAL PAGES ARE LISTED, THE WAY VISIONS LISTS THEM")
PAGES = [{"pos": 3, "domain": "complaintsboard.com",
          "url": "https://www.complaintsboard.com/ski-barn-c12"},
         {"pos": 7, "domain": "reddit.com",
          "url": "https://www.reddit.com/r/skiing/comments/abc/"}]
doc2 = Document(D.build_rep_proposal_docx(
    {"brand": "Ski Barn", "quote": QUOTE, "removal_pages": PAGES, "brackets": BRACKETS}))
t2 = [p.text for p in doc2.paragraphs if p.text.strip()]
check("the lead-in names the client",
      any(x.startswith("We have reviewed Ski Barn's online reputation") for x in t2), True)
check("every url is a bullet",
      [x for x in t2 if x.startswith("https://")], [p["url"] for p in PAGES])
check("it sits under Website Removals",
      t2.index("https://www.complaintsboard.com/ski-barn-c12") > t2.index("Website Removals"), True)
check("and before the rate card terms",
      t2.index("https://www.reddit.com/r/skiing/comments/abc/")
      < next(i for i, x in enumerate(t2) if x.startswith("Pricing for website removals")), True)

print("\nNO PAGES IS NO LEAD-IN")
doc3 = Document(D.build_rep_proposal_docx(
    {"brand": "Ski Barn", "quote": QUOTE, "brackets": BRACKETS}))
t3 = [p.text for p in doc3.paragraphs if p.text.strip()]
check("no dangling sentence",
      any(x.startswith("We have reviewed") for x in t3), False)
check("the section still builds", "Website Removals" in t3, True)
check("a domain with no url still lists",
      [x for x in [p.text for p in Document(D.build_rep_proposal_docx(
          {"brand": "Ski Barn", "quote": QUOTE, "brackets": BRACKETS,
           "removal_pages": [{"domain": "gripeo.com"}]})).paragraphs]
       if x.strip() == "gripeo.com"], ["gripeo.com"])

print("\nTHE SNAPSHOT THE PANEL DRAWS REACHES THE DOCUMENT")
# Negative-term demand, the tagged page-one results and the star split per
# location were rendered on screen and went no further. The pricing is an
# answer to those three tables.
SNAP = {"query": "sage dental reviews",
        "terms": [{"term": "sage dental reviews", "class": "watch", "volume": 720},
                  {"term": "sage dental lawsuit", "class": "negative", "volume": 170}],
        "organic": [{"pos": 1, "domain": "trustpilot.com", "owned": False,
                     "tactic": "site removal"},
                    {"pos": 3, "domain": "mysagedental.com", "owned": True,
                     "tactic": "owned \u2014 boost"},
                    {"pos": 6, "domain": "yelp.com", "owned": False,
                     "tactic": "suppression", "rating": 2.6}],
        "forums": [],
        "locations": [{"title": "Sage Dental of Midtown Atlanta", "profile_rating": 3.8,
                       "profile_reviews": 714, "neg_1": 54, "neg_2": 34, "weak_3": 17},
                      {"title": "Sage Dental of Tucker", "profile_rating": 3.9,
                       "profile_reviews": 671, "neg_1": 47, "neg_2": 26, "weak_3": 16}]}


def snap_doc(snapshot=SNAP, shots=None):
    return Document(D.build_rep_proposal_docx(
        {"brand": "Sage Dental", "quote": QUOTE, "snapshot": snapshot,
         "serp_shots": shots, "brackets": BRACKETS}))


sd = snap_doc()
st = [p.text for p in sd.paragraphs if p.text.strip()]
heads = [[c.text for c in t.rows[0].cells] for t in sd.tables]
check("it has its own heading", "Reputation Snapshot" in st, True)
check("negative terms", ["Term", "Class", "Searches/mo"] in heads, True)
check("page one with its routing", ["#", "Result", "Routing", "Rating"] in heads, True)
check("review profiles by location",
      ["Location", "Rating", "Reviews", "1\u2605", "2\u2605", "3\u2605"] in heads, True)
_terms = sd.tables[heads.index(["Term", "Class", "Searches/mo"])]
check("volume is formatted", _terms.rows[1].cells[2].text, "720")
_res = sd.tables[heads.index(["#", "Result", "Routing", "Rating"])]
check("owned is marked as owned", _res.rows[2].cells[2].text, "owned \u2014 owned \u2014 boost")
check("a third party is not", _res.rows[1].cells[2].text, "third party \u2014 site removal")
check("a missing rating is a dash", _res.rows[1].cells[3].text, "\u2014")
check("a rating shows", _res.rows[3].cells[3].text, "2.6\u2605")
check("the flagged total is stated",
      any(x.startswith("161 reviews sit at 1\u20132 stars") for x in st), True)

print("\nTHE SCREENSHOT SITS INSIDE IT, WITHOUT A SECOND HEADING")
sd2 = snap_doc(shots=[{"query": "sage dental reviews", "data_url": URL}])
st2 = [p.text for p in sd2.paragraphs if p.text.strip()]
check("image", len(sd2.inline_shapes), 1)
check("no duplicate Search Results heading", "Search Results" in st2, False)
check("the page-one sentence introduces it",
      any(x.startswith("Page one for") for x in st2), True)

print("\nNO SCAN IS NO SECTION")
check("nothing at all", "Reputation Snapshot" in
      [p.text for p in snap_doc(snapshot={}).paragraphs], False)
check("terms alone still draws it", "Reputation Snapshot" in
      [p.text for p in snap_doc(snapshot={"terms": SNAP["terms"]}).paragraphs], True)
check("a screenshot alone keeps its heading",
      "Search Results" in [p.text for p in snap_doc(
          snapshot={}, shots=[{"query": "q", "data_url": URL}]).paragraphs], True)

print("\n%d checks, %d failed" % (len(CHECKS), len(FAIL)))
if FAIL:
    print("FAILED: " + ", ".join(FAIL))
    sys.exit(1)
print("all ok")
