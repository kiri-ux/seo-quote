"""THREE THINGS THE DOCUMENT SAID THAT IT SHOULD NOT HAVE.

1. The Background section printed the client's own marketing line back at them
   -- "Create fast, professional CCTV drain survey reports with Drainify.
   Streamline inspections, simplify reporting, and impress clients". That text
   is the vocabulary sample the keyword build measures against; it was never
   copy for a document going to the client.

2. The authority paragraph opened "In practical terms that is roughly..." and
   the small-gap verdict pointed at "this proposal", both of which read as the
   document talking about itself.

3. Add-on markets were priced in the tool and absent from the document. The
   wording is Brendan's, from the TN Water & Air proposal -- the only one of the
   ten that carries them. (2026-09-09, Kiri)
"""
import importlib.util
import io
import os
import sys

os.environ.setdefault("DFS_LOGIN", "x")
os.environ.setdefault("DFS_PASSWORD", "x")
HERE = os.path.dirname(os.path.abspath(__file__))
SRCDIR = os.path.dirname(HERE)
sys.path.insert(0, SRCDIR)
SRC = os.path.join(SRCDIR, "app.py")
spec = importlib.util.spec_from_file_location("app", SRC)
app = importlib.util.module_from_spec(spec)
sys.modules["app"] = app
spec.loader.exec_module(app)

FAIL, CHECKS = [], []


def check(label, got, want):
    ok = got == want
    CHECKS.append(label)
    print(("  ok   " if ok else "  FAIL ") + label)
    if not ok:
        print("         got  %r\n         want %r" % (got, want))
        FAIL.append(label)


try:
    import docx
except ImportError:
    print("python-docx not installed here — cannot check the document")
    sys.exit(0)

TAGLINE = ("Create fast, professional CCTV drain survey reports with Drainify. "
           "Streamline inspections, simplify reporting, and impress clients—"
           "without the hassle.")
TERMS = ["sewer inspection software", "pacp software", "sewer crawler"]


def build(addons=1, client_rank=118):
    payload = {
        "brand": "Drainify", "business_desc": TAGLINE, "geo_values": [],
        "kw": {"all": [{"kw": t, "vol": 10} for t in TERMS],
               "ultra": [{"kw": t, "vol": 10} for t in TERMS],
               "competitive": [], "long_tail": [], "total_volume": 150},
        "table": [{"kw": t, "pos": 77} for t in TERMS],
        "signals": {"median_rival_rank": 214, "client_rank": client_rank,
                    "client_measured": True,
                    "rivals": [{"domain": "pipelogix.com", "rank": 180,
                                "appearances": 2}]},
        "pricing": {"client_tiers": {"base": 2950, "intermediate": 4000,
                                     "advanced": 5100},
                    "anchor": 1800, "markup_pct": 35, "addon_markets": addons,
                    "client_addon_per_market": {"base": 2655,
                                                "intermediate": 3600,
                                                "advanced": 4590},
                    "addon_discount_pct": 10, "min_term_months": 6,
                    "handoff": {}},
        "strategy": "Core SEO", "perf_on": False, "cpc": {}, "perf_extra": [],
        "perf_rows": [], "perf_override": {}, "practice_area": {},
        "state": "", "inputs": {"geo_values": []}, "serp": None}
    r = app.app.test_client().post("/api/proposal.docx", json=payload)
    assert r.status_code == 200, r.data[:300]
    doc = docx.Document(io.BytesIO(r.data))
    out = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    # The priced options are single-cell TABLES, not paragraphs — a border Word
    # and LibreOffice both draw. Read them too or half the document is invisible
    # to this test.
    for tb in doc.tables:
        for row in tb.rows:
            for cell in row.cells:
                for line in cell.text.split("\n"):
                    if line.strip():
                        out.append(line.strip())
    return out


txt = build()
blob = "\n".join(txt)

print("\nTHE CLIENT'S OWN MARKETING LINE IS NOT IN IT")
check("the tagline is gone", "Streamline inspections" in blob, False)
check("Background still introduces them",
      any(t.startswith("Drainify is requesting a proposal") for t in txt), True)
check("and still hands off to the recommendations",
      "Please find our recommendations and proposal below." in txt, True)

print("\nTHE AUTHORITY PARAGRAPH TALKS ABOUT THE CLIENT, NOT ITSELF")
check("no 'in practical terms'", "In practical terms" in blob, False)
check("no 'this proposal'", "in this proposal" in blob, False)
check("the measurement survives",
      "Their typical authority score is 214, against 118 for Drainify." in txt, True)
check("the verdict survives",
      any(t.startswith("That is a small gap.") for t in txt), True)
check("and the effort behind it",
      any("roughly 10 points of authority" in t and "40-80 referring domains" in t
          for t in txt), True)

print("\nADD-ON MARKETS ARE PRICED IN THE DOCUMENT")
check("the section is there", "Optional Add-On Markets" in txt, True)
check("with Brendan's lead-in",
      "In conjunction with the above SEO campaigns, add-on markets can be "
      "added as follows:" in txt, True)
for label, money in (("Base", "$2,655"), ("Intermediate", "$3,600"),
                     ("Advanced", "$4,590")):
    check(f"{label} rate",
          f"{label} Campaign: {money} per month per add-on market" in txt, True)
check("the market count is named",
      any("total of 1 additional market." in t for t in txt), True)
check("and it reads as singular", "1 additional markets" in blob, False)

print("\nTHE PLURAL, AND THE CASE WITH NO ADD-ONS")
t7 = "\n".join(build(addons=7))
check("seven reads as plural", "total of 7 additional markets." in t7, True)
t0 = "\n".join(build(addons=0))
check("no add-ons, no section", "Optional Add-On Markets" in t0, False)
check("the campaign options are still there", "Option 1: Base SEO Campaign" in t0, True)

print("\n%d checks, %d failed" % (len(CHECKS), len(FAIL)))
if FAIL:
    print("FAILED: " + ", ".join(FAIL))
    sys.exit(1)
print("all ok")
