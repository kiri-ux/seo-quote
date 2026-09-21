"""THE ORM DECK, FROM THE REPUTATION QUOTE.

The SEO row has had a deck since 2026-09-19; the reputation row had the letter
and nothing else, so the ORM slides were rebuilt by hand off the product deck
every time a partner presented one.

Checked here because a deck cannot be eyeballed in CI: the slide set per
strategy combination, the card row's geometry (it is sized to the number of
cards, so a Reactive-only quote must print one full-width card rather than one
card and three gaps), and every number against the handoff it was read off.
"""
import os
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.dirname(HERE))
import rep_pricing
import rep_pptx
from pptx import Presentation
from pptx.util import Emu, Inches

FAIL, RUN = [], []


def check(label, got, want):
    ok = got == want
    RUN.append(label)
    print(("  ok   " if ok else "  FAIL ") + label)
    if not ok:
        print("         got  %r\n         want %r" % (got, want))
        FAIL.append(label)


def quote(**kw):
    p = {"campaign": "bundle", "margin_pct": 0.35,
         "reviews": {"count": 0}, "articles": {"standard": 0},
         "search": {}, "shield": {}}
    p.update(kw)
    return rep_pricing.build_rep_quote(p)


SNAP = {
    "query": "sage dental reviews",
    "suggest": ["sage dental reviews florida", "sage dental care reviews",
                "sage dental ratings"],
    "organic": [
        {"pos": 1, "domain": "trustpilot.com", "tactic": "site removal"},
        {"pos": 3, "domain": "mysagedental.com", "owned": True,
         "tactic": "owned — boost"},
        {"pos": 6, "domain": "yelp.com", "tactic": "suppression", "rating": 2.6},
    ],
    "forums": [],
    "locations": [
        {"title": "Sage Dental of Midtown Atlanta", "profile_rating": 3.8,
         "profile_reviews": 714, "neg_1": 54, "neg_2": 34, "weak_3": 17},
        {"title": "Sage Dental of Tucker", "profile_rating": 3.9,
         "profile_reviews": 500, "neg_1": 47, "neg_2": 26, "weak_3": 16},
    ],
}


def deck(q, snap=None, **extra):
    d = {"brand": "Sage Dental", "order_no": "56310",
         "start_date": "2025-01-01", "end_date": "2025-03-30", "months": "3",
         "quote": q, "snapshot": snap if snap is not None else SNAP}
    d.update(extra)
    buf = rep_pptx.build_rep_proposal_pptx(d)
    return Presentation(buf), d


def titles(prs):
    """The heading on each slide, in order."""
    out = []
    for s in prs.slides:
        best = ""
        for sh in s.shapes:
            if not sh.has_text_frame:
                continue
            t = sh.text_frame.text.strip()
            if t.startswith("Online Reputation Management") and len(t) > len(best):
                best = t
        out.append(best or "(snapshot)")
    return out


def all_text(slide):
    out = []
    for sh in slide.shapes:
        if sh.has_text_frame:
            out.append(sh.text_frame.text)
        if getattr(sh, "has_table", False) and sh.has_table:
            for row in sh.table.rows:
                out += [c.text for c in row.cells]
    return "\n".join(out)


# ---------------------------------------------- what was sold is what prints
FULL = quote(reviews={"count": 10}, articles={"standard": 3},
             search={"bundle": True, "volume": 5000, "suppression": True,
                     "autosuggest": True, "term_sets": 1},
             shield={"locations": 3})
prs, d = deck(FULL)
check("the four-workstream quote prints five slides", len(prs.slides), 5)
check("and they are the product deck's own slides", titles(prs), [
    "Online Reputation Management",
    "Online Reputation Management - Strategy Details",
    "Online Reputation Management - Strategy Details",
    "Online Reputation Management - Reputation Snapshot",
    "Online Reputation Management Product Details"])
check("every strategy is on a strategy slide",
      all(s in "".join(all_text(x) for x in prs.slides)
          for s in ("Search Protection Bundle:",
                    "SEO Brand Shield & Asset Building:",
                    "Site/Article Removals:")), True)

# AS MANY STRATEGIES PER SLIDE AS FIT. The first cut gave each one a slide of
# its own, so a removals-only quote got five inches of white under one
# sentence. Reactive and Proactive together need 5.25in of the 5.48in a slide
# has, so they share one and the removals go to the second.
def strategy_slides(prs):
    return [s for s in prs.slides if "Strategy Details" in all_text(s)]


packed = strategy_slides(prs)
check("reactive and proactive share a slide",
      ("Search Protection Bundle:" in all_text(packed[0])
       and "SEO Brand Shield & Asset Building:" in all_text(packed[0])), True)
check("and the removals overflow to the next", len(packed), 2)

# ONE BOX PER STRATEGY. Reactive used to be a lead box, a floating bundle
# heading and three more boxes -- four cards for one product.
def boxes_below_heading(slide):
    return [sh for sh in slide.shapes
            if sh.has_text_frame and sh.top is not None
            and sh.top > Inches(1.1) and sh.text_frame.text.strip()]


check("reactive is drawn as one box",
      len([b for b in boxes_below_heading(packed[0])
           if "Search Protection Bundle:" in b.text_frame.text]), 1)
check("and that box holds every component",
      all(n + ":" in [b for b in boxes_below_heading(packed[0])
                      if "Search Protection Bundle:" in b.text_frame.text
                      ][0].text_frame.text
          for n, _ in rep_pptx.DOC["reactive_items"]), True)

# HER QUOTE: Review Removals + Reactive, which fit together.
HERS = quote(campaign="reactive", reviews={"count": 4},
             search={"bundle": True, "volume": 70, "suppression": True,
                     "autosuggest": True, "term_sets": 1})
hers, _ = deck(HERS)
check("review removals ride with reactive rather than taking a slide",
      len(strategy_slides(hers)), 1)

# A SLIDE WITH NOTHING BEHIND IT IS A SLIDE TO DELETE BY HAND.
REACTIVE = quote(campaign="reactive",
                 search={"bundle": True, "volume": 5000, "suppression": True,
                         "autosuggest": True, "term_sets": 1})
prs1, _ = deck(REACTIVE)
check("a reactive-only quote prints one strategy slide",
      [t for t in titles(prs1)].count(
          "Online Reputation Management - Strategy Details"), 1)
check("and says nothing about a brand shield",
      "SEO Brand Shield" in "".join(all_text(s) for s in prs1.slides), False)
check("and no removals slide",
      "Review Removals:" in "".join(all_text(s) for s in prs1.slides), False)
check("a scan with nothing in it prints no snapshot",
      "Reputation Snapshot" in " ".join(titles(deck(REACTIVE, snap={})[0])), False)
check("and a quote with no priced line prints no product details",
      titles(deck(quote(campaign="reactive"), snap={})[0]),
      ["Online Reputation Management"])

# ---------------------------------------------- the card row is sized to fit
def cards_of(prs):
    """The card rectangles on the Product Details slide: (left, width)."""
    slide = prs.slides[-1]
    out = []
    for sh in slide.shapes:
        if sh.top is not None and abs(sh.top - Inches(2.16)) < Emu(1000) \
           and sh.height and abs(sh.height - Inches(4.5)) < Emu(1000):
            out.append((sh.left, sh.width))
    return sorted(out)


def row_ok(cards, n):
    """n cards, left to right, inside the margins and never overlapping."""
    if len(cards) != n:
        return "got %d cards" % len(cards)
    gap = Inches(0.22)
    if cards[0][0] != Inches(0.55):
        return "first card does not start at the margin"
    right = cards[-1][0] + cards[-1][1]
    if abs(right - Inches(12.75)) > Emu(20000):
        return "row ends at %.3f in, not 12.75" % Emu(int(right)).inches
    for a, b in zip(cards, cards[1:]):
        if abs((b[0] - (a[0] + a[1])) - gap) > Emu(20000):
            return "cards are not one gap apart"
    return "ok"


check("four cards fill the row", row_ok(cards_of(prs), 4), "ok")
check("one card fills the row on its own", row_ok(cards_of(prs1), 1), "ok")
TWO = quote(search={"bundle": True, "volume": 5000, "suppression": True,
                    "autosuggest": True, "term_sets": 1},
            shield={"locations": 1000})
prs2, _ = deck(TWO)
check("two cards split it", row_ok(cards_of(prs2), 2), "ok")

# ---------------------------------------------- every number off the handoff
h = FULL["handoff"]
txt = all_text(prs.slides[-1])
check("the flight prints", "01/01/2025 - 03/30/2025" in txt, True)
check("months running prints", "Months Running: " in txt and "3" in txt, True)
check("the per-review RATE prints, not the order total",
      "$%s" % "{:,}".format(h["price_per_review_removal"]) in txt, True)
check("the per-page rate prints",
      "$%s" % "{:,}".format(h["price_per_standard_site_removal"]) in txt, True)
check("search protection's monthly prints",
      "$%s" % "{:,}".format(h["search_protection_monthly"]) in txt, True)
check("the brand shield's monthly prints",
      "$%s" % "{:,}".format(h["brand_shield_monthly"]) in txt, True)
# THE RECURRING MONEY OVER THE FLIGHT PLUS THE REMOVALS AT THEIR MAXIMUM.
# Pay on success, so it is a ceiling -- the same figure the row's tile is.
want_total = ((h["search_protection_monthly"] + h["brand_shield_monthly"]) * 3
              + h["reviews_count"] * h["price_per_review_removal"]
              + h["standard_sites"] * h["price_per_standard_site_removal"])
check("the total is the flight plus the removal maximums",
      "$%s" % "{:,}".format(want_total) in txt, True)
check("each bundle says what it scales with",
      ("5,000/mo measured" in txt, "3 locations measured" in txt), (True, True))

# ---------------------------------------------- the snapshot is the scan
# Found by its heading, not by its index: the slide count moves when the
# strategy sections pack differently.
snap_txt = all_text([s for s in prs.slides
                     if "Reputation Snapshot" in all_text(s)][0])
check("the query the suggestions came from prints",
      "sage dental reviews" in snap_txt, True)
check("the profiles print", "Sage Dental of Tucker" in snap_txt, True)
check("page one prints", "trustpilot.com" in snap_txt, True)
check("the tactic legend prints", "outrank it instead" in snap_txt, True)
# THE TILE IS A LINE ON THE QUOTE. 1 and 2 star across the ticked profiles --
# the same rows the flag count is read off, so they cannot disagree.
check("the removal tile counts 1 and 2 star", str(54 + 34 + 47 + 26) in snap_txt,
      True)
check("and the pie carries all three buckets",
      rep_pptx._stars(d), {"1": 101, "2": 60, "3": 33})

# ---------------------------------------------- the modifier is what's bold
# A legal suffix or an apostrophe must not decide whether anything is bold.
check("the brand splits off the modifier",
      rep_pptx._brand_split("sage dental reviews florida", "Sage Dental"),
      ("sage dental", " reviews florida"))
check("a suffix on the typed name does not stop it",
      rep_pptx._brand_split("cisney & o'donnell complaints",
                            "Cisney & O'Donnell PA"),
      ("cisney & o'donnell", " complaints"))
check("a phrase that does not lead with the brand is left whole",
      rep_pptx._brand_split("reviews for sage dental", "Sage Dental"),
      ("reviews for sage dental", ""))

# ---------------------------------------------- nothing runs off the slide
SLACK = Inches(1.5)
off = []
for n, s in enumerate(prs.slides, 1):
    for sh in s.shapes:
        if sh.left is None or sh.width is None:
            continue
        if (sh.left < -Inches(1.0) or sh.top < -Inches(0.7)
                or sh.left + sh.width > rep_pptx.SLIDE_W + SLACK
                or sh.top + sh.height > rep_pptx.SLIDE_H + SLACK):
            off.append((n, Emu(int(sh.left)).inches, Emu(int(sh.top)).inches))
check("every shape sits on the slide", off, [])

# THE LEGEND GOES UNDER THE TABLE, NOT THROUGH IT. PowerPoint grows a row to
# fit its text, so a legend placed at the estimated table height landed
# halfway up a table that had grown past it -- five results printed underneath.
wide = deck(FULL, snap=dict(
    SNAP, organic=[{"pos": i, "domain": "domain%d.com" % i,
                    "tactic": "suppression", "rating": 2.3}
                   for i in range(1, 13)]))[0]
snapshot = [s for s in wide.slides if "Reputation Snapshot" in all_text(s)][0]
mid = [sh for sh in snapshot.shapes
       if sh.left is not None and abs(sh.left - Inches(4.0)) < Inches(0.05)]
table = [sh for sh in mid if getattr(sh, "has_table", False) and sh.has_table]
legend = [sh for sh in mid if sh.has_text_frame
          and "outrank it instead" in sh.text_frame.text]
check("page one is capped so it cannot reach the legend",
      len(table) == 1 and len(table[0].table.rows) <= rep_pptx.MAX_RESULTS + 1,
      True)
check("and the legend starts below the table's last row",
      len(legend) == 1
      and legend[0].top >= table[0].top + table[0].height, True)
check("the rating column carries no vote count to wrap on",
      "(" in "".join(c.text for c in table[0].table.rows[1].cells), False)

# NOTHING UNDER THE TITLE. A navy corner block sat beneath the icon and the
# heading and read as a bar drawn through it.
det = prs.slides[-1]
banded = [sh for sh in det.shapes
          if sh.top is not None and sh.top < Inches(1.3)
          and sh.width is not None and sh.width < rep_pptx.SLIDE_W
          and not (sh.has_text_frame and sh.text_frame.text.strip())]
check("no blank shape overlaps the heading", banded, [])

print()
print("%d checks, %d failed" % (len(RUN), len(FAIL)))
print("all ok" if not FAIL else "FAILED: " + ", ".join(FAIL))
sys.exit(1 if FAIL else 0)
